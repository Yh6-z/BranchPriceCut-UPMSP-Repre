"""读取配置指定的 JD 算例数据，作为调度模型的输入。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from config import get_config


@dataclass(frozen=True, slots=True)
class SchedulingInstance:
    """无关并行机调度实例。加工时间矩阵采用“作业 × 机器”方向。"""

    num_machines: int
    num_jobs: int
    instance_id: int
    job_ids: tuple[int, ...]
    machine_ids: tuple[int, ...]
    weights: tuple[int, ...]
    processing_times: tuple[tuple[float, ...], ...]

    def weight(self, job_id: int) -> int:
        return self.weights[job_id - 1]

    def processing_time(self, job_id: int, machine_id: int) -> float:
        return self.processing_times[job_id - 1][machine_id - 1]


def build_instance_path(configuration: object) -> Path:
    """根据规模和精度定位 Excel 实例。

    输入：含机器数、作业数、实例号及可选数据目录/精度的配置 object。
    输出：实例文件的 pathlib.Path。
    """

    filename = (
        f"{configuration.num_machine}_{configuration.num_job}_"
        f"TWCT_{configuration.instance_id}.xlsx"
    )
    directory = getattr(configuration, "data_dir", Path(__file__).resolve().parents[2] / "data" / "JD" / "JD Instances")
    decimals = getattr(configuration, "processing_time_decimals", 1)
    folder = "JD_Instances_Processing_times_1_decimal" if decimals == 1 else "JD_Instances_Processing_times_2_decimals"
    return Path(directory) / folder / filename


def load_instance(configuration: object | None = None) -> SchedulingInstance:
    """顺序读取权重和加工时间，再转换为作业 × 机器矩阵。

    输入：含规模、实例编号、数据路径及精度的配置 object 或 None。
    输出：SchedulingInstance；权重为 tuple[int, ...]，加工时间为
    tuple[tuple[float, ...], ...]，外层为作业，内层为机器。
    """

    config = get_config([]) if configuration is None else configuration
    workbook = load_workbook(build_instance_path(config), read_only=True, data_only=True)
    weights_sheet = workbook["Weights"]
    processing_times_sheet = workbook["Processing times"]

    job_ids = tuple(range(1, config.num_job + 1))
    machine_ids = tuple(range(1, config.num_machine + 1))
    weights = tuple(
        int(row[0]) for row in weights_sheet.iter_rows(
            min_row=2, max_row=config.num_job + 1, min_col=2, max_col=2, values_only=True,
        )
    )
    # 只读工作表逐单元格访问会反复解析 XML；各表顺序扫描一次即可。
    machine_rows = tuple(processing_times_sheet.iter_rows(
        min_row=2, max_row=config.num_machine + 1,
        min_col=2, max_col=config.num_job + 1, values_only=True,
    ))
    processing_times = tuple(
        tuple(float(machine_rows[machine_id - 1][job_id - 1]) for machine_id in machine_ids)
        for job_id in job_ids
    )
    workbook.close()

    return SchedulingInstance(
        num_machines=config.num_machine,
        num_jobs=config.num_job,
        instance_id=config.instance_id,
        job_ids=job_ids,
        machine_ids=machine_ids,
        weights=weights,
        processing_times=processing_times,
    )


if __name__ == "__main__":
    instance = load_instance(get_config())
    print(f"machines={instance.num_machines}")
    print(f"jobs={instance.num_jobs}")
    print(f"instance_id={instance.instance_id}")
    print(f"weights={len(instance.weights)}")
    print(f"processing_time_shape=({instance.num_jobs}, {instance.num_machines})")
